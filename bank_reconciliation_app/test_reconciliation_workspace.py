import ast
from datetime import datetime
from io import BytesIO
from pathlib import Path
import re
import types
import unittest

import pandas as pd
from openpyxl import load_workbook

from reconciliation_workspace import ReconciliationWorkspace, REQUIRED


def sample():
    ledger = pd.DataFrame({
        "Data": [datetime(2026, 1, 1), datetime(2026, 1, 2)],
        "Descrição": ["Pagamento, Abc!", "Outro"], "Valor": [100, 50],
        "Referencia": ["000123", "NA"], "Em branco": ["", ""],
        "Ver": ["original", "preservar"], "ordem": ["cliente-a", "cliente-b"],
        "nota": ["=texto", "nota original"], "index": ["001", "002"],
        "DateDiff": [7, 8],
    })
    bank = pd.DataFrame({
        "Data": [datetime(2026, 1, 1), datetime(2026, 1, 2)],
        "Descrição": ["Pagamento Abc", "Outro"], "Valor": [100, 50],
        "BancoRef": ["000001", "000002"],
    })
    return {"CNT": ledger, "BANCO": bank}


def reopen(workspace):
    return ReconciliationWorkspace.read(BytesIO(workspace.export()))


class WorkbookTests(unittest.TestCase):
    def test_roundtrip_preserves_source_columns_values_and_ids(self):
        workspace = ReconciliationWorkspace.from_sheets(sample())
        workspace.settings = {"saldo_i_cnt": 100.25, "saldo_i_banco": 99, "match_total": 100, "match_parcial": 80}
        original_ids = {side: list(frame.index) for side, frame in workspace.sources.items()}
        order = workspace.reconcile([original_ids["CNT"][0]], [original_ids["BANCO"][0]], "Conferido")
        self.assertEqual(order, 1)
        for _ in range(3):
            workspace = reopen(workspace)
            self.assertEqual(workspace.next_id, 2)
            self.assertEqual(len(workspace.history), 2)
            self.assertEqual(workspace.settings["saldo_i_cnt"], 100.25)
            for side, expected in sample().items():
                actual = workspace.sources[side].loc[original_ids[side]].reset_index(drop=True)
                pd.testing.assert_frame_equal(actual, expected, check_dtype=False)
            self.assertEqual(workspace.matches[1]["nota"], "Conferido")
        self.assertEqual(workspace.reconcile(workspace.pending["CNT"], workspace.pending["BANCO"]), 2)

    def test_history_survives_undo_and_id_is_never_reused(self):
        workspace = ReconciliationWorkspace.from_sheets(sample())
        workspace.reconcile(workspace.pending["CNT"], workspace.pending["BANCO"], "Primeira")
        workspace = reopen(workspace)
        workspace.undo([1])
        workspace = reopen(workspace)
        self.assertEqual(workspace.matches, {})
        self.assertEqual(workspace.next_id, 2)
        self.assertEqual([e["Acao"] for e in workspace.history], ["Manual"] * 4 + ["Anulada"] * 4)
        self.assertEqual(workspace.reconcile(workspace.pending["CNT"], workspace.pending["BANCO"]), 2)
        self.assertEqual(len(reopen(workspace).history), 12)

    def test_legacy_history_uses_maximum_across_both_sides(self):
        sheets = {side: pd.DataFrame(columns=REQUIRED) for side in ("CNT", "BANCO")}
        sheets["MATCH CNT"] = pd.DataFrame([[datetime(2026, 1, 1), "a", 0, "120", "cnt"]], columns=REQUIRED + ["ordem", "nota"])
        sheets["MATCH BANCO"] = pd.DataFrame([[datetime(2026, 1, 1), "b", 0, 125, "banco"]], columns=REQUIRED + ["ordem", "nota"])
        workspace = reopen(ReconciliationWorkspace.from_sheets(sheets))
        self.assertEqual(workspace.next_id, 126)
        self.assertEqual(set(workspace.matches), {120, 125})
        self.assertEqual([e["Acao"] for e in workspace.history], ["Importada", "Importada"])
        workspace.undo([125])
        self.assertEqual(len(workspace.pending["BANCO"]), 1)
        self.assertEqual(workspace.next_id, 126)

    def test_empty_and_no_new_matches_can_be_saved(self):
        for sheets in (sample(), {side: pd.DataFrame(columns=REQUIRED + ["Vazia"]) for side in ("CNT", "BANCO")}):
            workspace = reopen(ReconciliationWorkspace.from_sheets(sheets))
            self.assertEqual(workspace.next_id, 1)
            self.assertEqual(workspace.matches, {})
            self.assertEqual(len(workspace.pending["CNT"]), len(sheets["CNT"]))

    def test_failed_and_empty_matches_do_not_consume_ids(self):
        workspace = ReconciliationWorkspace.from_sheets(sample())
        for ledger, bank in (([], []), ([workspace.pending["CNT"][0]], [workspace.pending["BANCO"][1]]),
                             ([workspace.pending["CNT"][0]] * 2, [])):
            with self.assertRaises(ValueError):
                workspace.reconcile(ledger, bank)
        self.assertEqual(workspace.next_id, 1)
        self.assertEqual(workspace.history, [])

    def test_appended_movements_and_extra_columns(self):
        workspace = ReconciliationWorkspace.from_sheets(sample())
        data = pd.read_excel(BytesIO(workspace.export()), sheet_name=None, dtype=object, keep_default_na=False)
        data["CNT"]["Nova coluna"] = ["x", "y"]
        added = {c: "" for c in data["CNT"].columns}
        added.update({"Data": datetime(2026, 1, 3), "Descrição": "Novo", "Valor": 1, "Nova coluna": "z"})
        data["CNT"] = pd.concat([data["CNT"], pd.DataFrame([added])], ignore_index=True)
        workspace = reopen(ReconciliationWorkspace.from_sheets(data))
        self.assertEqual(len(workspace.pending["CNT"]), 3)
        self.assertEqual(workspace.sources["CNT"]["Nova coluna"].tolist(), ["x", "y", "z"])

    def test_invalid_legacy_ids_and_partial_history_rejected(self):
        for value in ("abc", "", -1, 1.5):
            sheets = sample()
            sheets["MATCH CNT"] = pd.DataFrame([[datetime(2026, 1, 1), "a", 1, value]], columns=REQUIRED + ["ordem"])
            sheets["MATCH BANCO"] = pd.DataFrame(columns=REQUIRED + ["ordem"])
            with self.assertRaises(ValueError):
                ReconciliationWorkspace.from_sheets(sheets)
        sheets = sample()
        sheets["MATCH CNT"] = pd.DataFrame()
        with self.assertRaises(ValueError):
            ReconciliationWorkspace.from_sheets(sheets)

    def test_duplicate_persistent_ids_rejected(self):
        workspace = ReconciliationWorkspace.from_sheets(sample())
        sheets = pd.read_excel(BytesIO(workspace.export()), sheet_name=None, dtype=object, keep_default_na=False)
        column = workspace.metadata["CNT"]["row_id"]
        sheets["CNT"].loc[1, column] = sheets["CNT"].loc[0, column]
        with self.assertRaises(ValueError):
            ReconciliationWorkspace.from_sheets(sheets)

    def test_text_is_not_converted_to_formula(self):
        workspace = ReconciliationWorkspace.from_sheets(sample())
        book = load_workbook(BytesIO(workspace.export()))
        column = list(workspace.sources["CNT"].columns).index("nota") + 1
        self.assertEqual(book["CNT"].cell(2, column).value, "=texto")
        self.assertEqual(book["CNT"].cell(2, column).data_type, "s")

    def test_one_sided_offsetting_movements_can_be_undone(self):
        sheets = sample()
        sheets["CNT"]["Valor"] = [100, -100]
        workspace = ReconciliationWorkspace.from_sheets(sheets)
        workspace.reconcile(workspace.pending["CNT"], [], "Estorno")
        workspace = reopen(workspace)
        workspace.undo([1])
        self.assertEqual(len(workspace.pending["CNT"]), 2)
        self.assertEqual(len(workspace.pending["BANCO"]), 2)
        self.assertEqual(reopen(workspace).next_id, 2)

    def test_legacy_per_row_notes_preserved(self):
        sheets = {side: pd.DataFrame(columns=REQUIRED) for side in ("CNT", "BANCO")}
        sheets["MATCH CNT"] = pd.DataFrame([[datetime(2026, 1, 1), "a", 1, 5, "nota 1"], [datetime(2026, 1, 1), "b", 1, 5, "nota 2"]], columns=REQUIRED + ["ordem", "nota"])
        sheets["MATCH BANCO"] = pd.DataFrame([[datetime(2026, 1, 1), "c", 2, 5, "banco"]], columns=REQUIRED + ["ordem", "nota"])
        workspace = reopen(ReconciliationWorkspace.from_sheets(sheets))
        result = pd.read_excel(BytesIO(workspace.export()), sheet_name="MATCH CNT")
        self.assertEqual(result["nota"].tolist(), ["nota 1", "nota 2"])


class AppMatchingTests(unittest.TestCase):
    def test_best_candidate_is_the_row_saved_and_exported(self):
        tree = ast.parse(Path("reconciliacoes.py").read_text(encoding="utf-8"))
        functions = [n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in {"match_transactions", "normalize_description"}]
        from fuzzywuzzy import fuzz
        sheets = sample()
        sheets["BANCO"].loc[1, "Valor"] = 100
        sheets["BANCO"].loc[1, "Descrição"] = "Unrelated"
        workspace = ReconciliationWorkspace.from_sheets(sheets)
        namespace = {"re": re, "fuzz": fuzz, "st": types.SimpleNamespace(session_state=types.SimpleNamespace(workspace=workspace))}
        exec(compile(ast.Module(body=functions, type_ignores=[]), "app-matching", "exec"), namespace)
        selected_id = workspace.pending["BANCO"][0]
        remaining_id = workspace.pending["BANCO"][1]
        namespace["match_transactions"](workspace.view("CNT"), workspace.view("BANCO"))
        self.assertEqual(workspace.matches[1]["BANCO"], [selected_id])
        self.assertEqual(workspace.pending["BANCO"], [remaining_id])
        workspace = reopen(workspace)
        self.assertEqual(workspace.sources["BANCO"].loc[selected_id, "BancoRef"], "000001")
        self.assertEqual(workspace.next_id, 2)


if __name__ == "__main__":
    unittest.main()
